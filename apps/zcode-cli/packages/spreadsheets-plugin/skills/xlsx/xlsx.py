#!/usr/bin/env python3
"""xlsx.py — the workbook toolbox.

Original work for this repository, written against the public APIs of openpyxl
(MIT), XlsxWriter (BSD-2-Clause) and the LibreOffice headless CLI; see the
plugin NOTICE.md. No third-party source is vendored.

Subcommands:

    init      create a workbook from a small spec (assumptions + one table)
    inspect   print a workbook's sheets, dimensions, formulas and formats
    recalc    recalculate with LibreOffice headless (delegates to
              scripts/recalc.py so there is exactly one recalc path)
    audit     structural checks from quality/pipeline.md (merged cells in
              data regions, formula errors, colour-role violations)

Usage:

    python3 xlsx.py init model.xlsx --assumption growth=0.05 --header region --header revenue
    python3 xlsx.py inspect model.xlsx
    python3 xlsx.py audit model.xlsx

Requires: Python >= 3.10. `init` and `audit` need openpyxl; `recalc` needs
LibreOffice; `inspect` works on any xlsx openpyxl can read.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

# Colour roles from engines/design.md §1, as openpyxl ARGB.
ROLE_BLUE = "FF0000FF"      # input
ROLE_BLACK = "FF000000"     # formula
ROLE_GREEN = "FF008000"     # cross-sheet link
ROLE_RED = "FFFF0000"       # external link
ROLE_YELLOW = "FFFFFF00"    # assumption fill

FORMULA_ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#NULL!", "#NUM!", "#N/A")


# --------------------------------------------------------------------- helpers


def _load_openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ImportError as error:  # pragma: no cover - environment dependent
        raise SystemExit(
            "openpyxl is required for this command: python3 -m pip install openpyxl"
        ) from error
    return openpyxl


def _parse_assumptions(pairs: list[str]) -> list[tuple[str, float]]:
    assumptions: list[tuple[str, float]] = []
    for pair in pairs:
        if "=" not in pair:
            raise SystemExit(f"assumption must be name=value, got: {pair}")
        name, raw = pair.split("=", 1)
        try:
            assumptions.append((name.strip(), float(raw)))
        except ValueError as error:
            raise SystemExit(f"assumption {name!r} is not a number: {raw!r}") from error
    return assumptions


# ---------------------------------------------------------------------- init


def cmd_init(args: argparse.Namespace) -> int:
    """Create a workbook: an Assumptions sheet and one data sheet."""
    openpyxl = _load_openpyxl()
    from openpyxl.styles import Font, PatternFill

    assumptions = _parse_assumptions(args.assumption)
    if not args.header:
        raise SystemExit("init needs at least one --header")

    workbook = openpyxl.Workbook()
    inputs = workbook.active
    inputs.title = "Assumptions"

    blue = Font(color=ROLE_BLUE)
    yellow = PatternFill(fill_type="solid", start_color=ROLE_YELLOW, end_color=ROLE_YELLOW)
    black = Font(color=ROLE_BLACK)

    for row, (name, value) in enumerate(assumptions, start=1):
        inputs.cell(row=row, column=1, value=name)
        cell = inputs.cell(row=row, column=2, value=value)
        cell.font = blue
        cell.fill = yellow
    inputs.column_dimensions["A"].width = 24
    inputs.column_dimensions["B"].width = 14

    data = workbook.create_sheet("Data")
    for col, header in enumerate(args.header, start=1):
        header_cell = data.cell(row=1, column=col, value=header)
        header_cell.font = Font(bold=True, color=ROLE_BLACK)
    data.freeze_panes = "A2"
    for col in range(1, len(args.header) + 1):
        data.column_dimensions[data.cell(row=1, column=col).column_letter].width = 16

    workbook.save(args.file)
    print(json.dumps({"created": args.file, "sheets": ["Assumptions", "Data"],
                      "assumptions": [name for name, _ in assumptions]}, indent=2))
    return 0


# ------------------------------------------------------------------- inspect


def cmd_inspect(args: argparse.Namespace) -> int:
    openpyxl = _load_openpyxl()
    workbook = openpyxl.load_workbook(args.file, data_only=False)
    report: dict[str, object] = {"file": args.file, "sheets": []}
    for sheet in workbook.worksheets:
        formulas = 0
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas += 1
        report["sheets"].append(
            {
                "name": sheet.title,
                "state": sheet.sheet_state,
                "dimensions": sheet.dimensions,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "formulas": formulas,
                "merged": len(sheet.merged_cells.ranges),
                "freeze": sheet.freeze_panes,
            }
        )
    print(json.dumps(report, indent=2))
    return 0


# --------------------------------------------------------------------- audit


def cmd_audit(args: argparse.Namespace) -> int:
    """The mechanical part of quality/pipeline.md: errors, merges, colour roles."""
    openpyxl = _load_openpyxl()
    workbook = openpyxl.load_workbook(args.file, data_only=True)
    findings: list[dict[str, object]] = []

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip() in FORMULA_ERRORS:
                    findings.append(
                        {"sheet": sheet.title, "cell": cell.coordinate,
                         "code": "formula_error", "value": cell.value.strip()}
                    )
        for merged in sheet.merged_cells.ranges:
            # A merge is a defect inside a data region (row >= 2), fine in a
            # title band (row 1).
            if merged.min_row >= 2:
                findings.append(
                    {"sheet": sheet.title, "cell": str(merged),
                     "code": "merged_in_data_region", "value": None}
                )

    # Colour roles: a formula must not be blue (an input colour).
    formula_workbook = openpyxl.load_workbook(args.file, data_only=False)
    for sheet in formula_workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                    continue
                color = getattr(cell.font, "color", None)
                rgb = getattr(color, "rgb", None)
                if isinstance(rgb, str) and rgb.upper() == ROLE_BLUE:
                    findings.append(
                        {"sheet": sheet.title, "cell": cell.coordinate,
                         "code": "formula_in_input_colour", "value": cell.value}
                    )

    verdict = "fail" if any(f["code"] == "formula_error" for f in findings) else "pass"
    print(json.dumps({"file": args.file, "verdict": verdict, "findings": findings}, indent=2))
    return 0 if verdict == "pass" else 1


# -------------------------------------------------------------------- recalc


def cmd_recalc(args: argparse.Namespace) -> int:
    """Delegate to scripts/recalc.py — one recalculation path, not two."""
    import subprocess
    import sys as _sys
    from pathlib import Path

    recalc = Path(__file__).resolve().parent / "scripts" / "recalc.py"
    if not recalc.exists():
        raise SystemExit(f"recalc.py not found next to this script: {recalc}")
    result = subprocess.run([_sys.executable, str(recalc), args.file])
    return result.returncode


# ----------------------------------------------------------------------- cli


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="xlsx.py", description=__doc__.splitlines()[0])
    sub = parser.add_subparsers(dest="command", required=True)

    init = sub.add_parser("init", help="create a workbook from a small spec")
    init.add_argument("file")
    init.add_argument("--assumption", action="append", default=[],
                      help="name=value, repeatable; lands in the Assumptions sheet")
    init.add_argument("--header", action="append", default=[],
                      help="one column header for the Data sheet, repeatable")
    init.set_defaults(func=cmd_init)

    inspect = sub.add_parser("inspect", help="print sheets, dimensions and formula counts")
    inspect.add_argument("file")
    inspect.set_defaults(func=cmd_inspect)

    audit = sub.add_parser("audit", help="structural checks from quality/pipeline.md")
    audit.add_argument("file")
    audit.set_defaults(func=cmd_audit)

    recalc = sub.add_parser("recalc", help="recalculate with LibreOffice (delegates to recalc.py)")
    recalc.add_argument("file")
    recalc.set_defaults(func=cmd_recalc)

    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    raise SystemExit(main())
